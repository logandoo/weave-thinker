# Copyright (c) 2026 Weave Thinker Contributors
# SPDX-License-Identifier: Apache-2.0

"""Office → PDF / HTML conversion service (2026-09-19, HTML 2026-09-21).

Server-side rendering path for Office previews (reference: deepseek-harness
``dsh-office-to-pdf`` — LibreOffice + client PDF rendering). The Node kit is
not portable to this Python stack, so the equivalent is a headless
``soffice`` subprocess:

* content-version cache key ``sha256(path:mtime_ns:size:ext)`` — a changed
  source re-converts, repeated previews are served from disk;
* one conversion at a time (``Semaphore(1)``), per-run isolated
  ``UserInstallation`` profile so concurrent soffice instances cannot fight
  over the user profile lock;
* bounded cache (LRU by mtime) and hard timeout with process-group kill.

Spreadsheets can additionally be rendered to a self-contained HTML document
(``convert_to_html``): the Calc HTML export keeps each sheet as one
natural-width table (no print pagination), preserves styles/merges, and embeds
charts as images which are inlined as data URIs and sanitized through a strict
whitelist before the browser ever sees them. This is the preview path for
``.xlsx``/``.xls``/``.ods``; word/ppt keep the PDF path.

When no soffice binary is available the caller maps
:class:`OfficePreviewUnavailable`/``None`` to HTTP 501 and the frontend falls
back to its client-side renderers — office preview never becomes unavailable
just because the host lacks LibreOffice.
"""
from __future__ import annotations

import base64
import hashlib
import html as html_lib
import logging
import os
import re
import shutil
import signal
import subprocess
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from html.parser import HTMLParser
from pathlib import Path
from typing import Callable
from urllib.parse import unquote

logger = logging.getLogger(__name__)

SUPPORTED_OFFICE_EXT = frozenset({
    ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx",
    ".odt", ".ods", ".odp", ".rtf",
})

# Spreadsheet formats that get the HTML preview path (2026-09-21): LibreOffice
# print-to-PDF pagination splits wide tables across pages and cuts long cell
# text at page boundaries; the Calc HTML export keeps every sheet as one
# natural-width table (horizontal scroll + CSS wrapping in the frontend) and
# embeds charts as images.
SPREADSHEET_EXT = frozenset({".xls", ".xlsx", ".ods"})

# Spreadsheet formats whose sheets may carry charts. LibreOffice's default
# paginated PDF export clips charts at page boundaries (user report
# 2026-09-19: sales_dashboard.xlsx lost chart categories / split the pie).
#
# Two mechanisms, chosen per format:
# * .xlsx → print-setup patch (fitToPage) in a spool copy + plain ``pdf`` —
#   works on every LibreOffice version (verified on prod LO 7.3, which
#   silently ignores the JSON filter options introduced in 7.4);
# * .xls / .ods → ``SinglePageSheets`` JSON filter (LO >= 7.4) with the
#   paginated fallback below for tall sheets / ignored options.
_CALC_SINGLE_PAGE_EXT = frozenset({".xls", ".ods"})
_XLSX_EXT = ".xlsx"

# Heuristic guard: a sheet taller than this many rows is exported paginated
# (fit-to-one-page would shrink it to unreadable scale).
_XLSX_FIT_MAX_ROWS = 200

# Bump whenever the conversion pipeline changes so stale cached PDFs are
# not served under an unchanged source (v4 = xlsx print-setup patch;
# v5 = patch hardening: paired pageSetUpPr, sheetProtection prefix,
# well-formedness validation; v6 = spreadsheet HTML preview path;
# v7 = workbook column-width injection + multi-colgroup stripping;
# v8 = sanitizer marker + CSS escape guard + early size budgets).
_CACHE_VERSION = "8"

# --- Spreadsheet HTML preview (2026-09-21) ---------------------------------
# Sanitizer/output guards: the converted document is served to a sandboxed
# iframe, and the sanitizer strips anything active or off-disk before it.
_HTML_MAX_BYTES = 16 * 1024 * 1024
_HTML_ASSET_MAX_BYTES = 8 * 1024 * 1024
_HTML_ASSET_EXT = frozenset({".png", ".gif", ".jpg", ".jpeg", ".bmp", ".webp"})
_HTML_ASSET_MIME = {
    ".png": "image/png",
    ".gif": "image/gif",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".bmp": "image/bmp",
    ".webp": "image/webp",
}
# Element/attribute whitelist mirrors what LibreOffice's Calc HTML export
# emits; everything else (scripts, frames, forms, on* handlers, data-* payload
# blobs, external URLs) is dropped.
_HTML_ALLOWED_TAGS = frozenset({
    "html", "head", "style", "body",
    "hr", "p", "center", "h1", "h2", "h3", "h4", "h5", "h6",
    "table", "colgroup", "col", "thead", "tbody", "tfoot", "tr", "td", "th",
    "font", "b", "i", "u", "em", "strong", "s", "strike", "sub", "sup",
    "small", "big", "span", "div", "br", "a", "img",
})
_HTML_VOID_TAGS = frozenset({"br", "hr", "col", "img"})
# Bare structural tags are rebuilt by the sanitizer (own doctype/head/body).
_HTML_SKELETON_TAGS = frozenset({"html", "head", "body"})
# Text content of these never reaches the preview.
_HTML_SUPPRESS_TAGS = frozenset({"script", "noscript", "textarea", "template", "title"})
_HTML_ALLOWED_ATTRS = frozenset({
    "align", "valign", "width", "height", "colspan", "rowspan", "border",
    "cellspacing", "cellpadding", "bgcolor", "color", "face", "size",
    "title", "alt", "nowrap", "name", "span",
})
# CSS is not tokenized here, so any backslash is rejected outright: CSS escape
# sequences (`u\72l(...)`, `@\69 mport`) would otherwise bypass the substring
# checks below (A4.9 2026-09-21). LibreOffice's own stylesheet contains none.
_HTML_CSS_FORBIDDEN = ("url(", "@import", "expression(", "javascript:", "behavior:", "\\")
# Marker proving the artifact passed the sanitizer (cache trust check).
_HTML_PREVIEW_MARKER = 'data-wt-preview="1"'
# Memory bounds: raw export is size-checked BEFORE parsing, inlined assets get
# a running byte budget, and the assembled document must fit _HTML_MAX_BYTES.
_HTML_MAX_RAW_CHARS = 48 * 1024 * 1024
_HTML_ASSET_TOTAL_MAX_BYTES = 24 * 1024 * 1024
# openpyxl width extraction decompresses the whole archive; refuse oversized
# expansion (a 30MB deflated workbook can inflate into millions of cells).
_HTML_WIDTHS_MAX_UNCOMPRESSED_BYTES = 128 * 1024 * 1024
# Readability layer injected after the LibreOffice stylesheet: text wraps
# inside cells (no clipping) and wide tables scroll horizontally instead of
# being split across pages.
_HTML_PREVIEW_CSS = (
    "html,body{margin:0;padding:0;background:#fff;}"
    "body{padding:10px 12px 28px;}"
    "table{border-collapse:collapse;}"
    "td,th{overflow-wrap:anywhere;}"
    "img{max-width:100%;height:auto;}"
)

# Single-page mode flattens a whole sheet onto one page. Beyond this page
# height (PDF points; A4 is ~842pt) the page is unreadably large and can even
# exceed the browser's max canvas dimensions, so such workbooks fall back to
# the paginated export (charts may clip, tables stay readable) — A4.9 I-2.
_SINGLE_PAGE_MAX_PT = 4000
_MEDIABOX_RE = re.compile(rb"/MediaBox\s*\[([^\]]+)\]")

_COMMON_SOFFICE_PATHS = (
    "/usr/bin/soffice",
    "/usr/local/bin/soffice",
    "/opt/libreoffice/program/soffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
)

_SEM = None


class OfficePreviewError(RuntimeError):
    """Base error for office preview conversion."""


class OfficePreviewUnavailable(OfficePreviewError):
    """No usable soffice binary is available."""


class OfficePreviewFailed(OfficePreviewError):
    """soffice was found but the conversion failed or timed out."""


def _conversion_semaphore() -> "asyncio.Semaphore":  # noqa: F821 - lazy import
    global _SEM
    if _SEM is None:
        import asyncio

        _SEM = asyncio.Semaphore(1)
    return _SEM


def is_supported(name: str) -> bool:
    return Path(name).suffix.lower() in SUPPORTED_OFFICE_EXT


def is_spreadsheet(name: str) -> bool:
    return Path(name).suffix.lower() in SPREADSHEET_EXT


def _is_valid_html(path: Path) -> bool:
    """HTML cache entries are trusted only after a marker check.

    The marker is emitted by :func:`sanitize_office_html`; a planted/poisoned
    file in the shared cache dir can therefore never be served as a preview
    (same threat model as the ``%PDF-`` magic check).
    """
    try:
        with open(path, "rb") as handle:
            head = handle.read(1024).lstrip().lower()
        return head.startswith(b"<!doctype html") and _HTML_PREVIEW_MARKER.encode() in head
    except OSError:
        return False


# Reading column widths needs a non-read-only openpyxl load; keep it bounded so
# a huge uploaded workbook cannot stall the worker.
_HTML_WIDTHS_MAX_SRC_BYTES = 32 * 1024 * 1024
_HTML_DEFAULT_COL_WIDTH_PX = round(8.43 * 7 + 5)


def _sheet_column_widths(src: Path) -> list[tuple[str, list[int]]] | None:
    """Per-sheet column widths in px, in workbook order (all sheet states).

    LibreOffice's Calc HTML export only keeps the first column's width, so the
    browser auto-fits (and squeezes) every table. The preview re-injects the
    workbook's intended widths; ``None`` when openpyxl cannot read the file
    (the HTML then keeps the browser's auto layout).
    """
    try:
        if src.stat().st_size > _HTML_WIDTHS_MAX_SRC_BYTES:
            return None
        with zipfile.ZipFile(src) as archive:
            uncompressed = sum(info.file_size for info in archive.infolist())
        if uncompressed > _HTML_WIDTHS_MAX_UNCOMPRESSED_BYTES:
            return None
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except (OSError, ImportError, zipfile.BadZipFile):
        return None
    try:
        wb = load_workbook(src, read_only=False, data_only=True)
    except Exception:  # noqa: BLE001 - width metadata is best-effort
        return None
    try:
        plan: list[tuple[str, list[int]]] = []
        for ws in wb.worksheets:
            widths: list[int] = []
            for col in range(1, min(int(ws.max_column or 0), 512) + 1):
                dim = ws.column_dimensions.get(get_column_letter(col))
                width = dim.width if dim is not None and dim.width else None
                widths.append(
                    round(width * 7 + 5) if width else _HTML_DEFAULT_COL_WIDTH_PX
                )
            plan.append((ws.title, widths))
        return plan
    except Exception:  # noqa: BLE001
        return None
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass


_HTML_TABLE_RE = re.compile(r"<table\b[^>]*>(.*?)</table>", re.S | re.I)
_HTML_ROW_RE = re.compile(r"<tr\b[^>]*>(.*?)</tr>", re.S | re.I)
_HTML_CELL_RE = re.compile(r"<t[dh]\b([^>]*)>", re.I)
_HTML_COLSPAN_RE = re.compile(r"colspan\s*=\s*[\"']?(\d+)", re.I)
_HTML_SHEET_HEADING_RE = re.compile(
    r"<h1>\s*Sheet\s+\d+\s*:\s*<em>(.*?)</em>\s*</h1>", re.S | re.I
)


def _html_table_column_counts(html: str) -> list[int]:
    """Column count per ``<table>`` (colspan-aware metadata scan)."""
    counts: list[int] = []
    for table in _HTML_TABLE_RE.finditer(html):
        max_cols = 0
        for row in _HTML_ROW_RE.finditer(table.group(1)):
            n = 0
            for cell in _HTML_CELL_RE.finditer(row.group(1)):
                span = _HTML_COLSPAN_RE.search(cell.group(1))
                n += int(span.group(1)) if span else 1
            max_cols = max(max_cols, n)
        counts.append(max_cols)
    return counts


def _html_table_width_plan(
    html: str, sheets: list[tuple[str, list[int]]]
) -> list[list[int]]:
    """Per-table column widths, aligned by the export's sheet headings.

    The Calc HTML export has one ``<h1>Sheet N: <em>NAME</em></h1>`` heading
    per exported sheet; a table is only width-injected when its sheet can be
    resolved and the source has at least as many columns as the table uses.
    """
    counts = _html_table_column_counts(html)
    headings = [
        html_lib.unescape(m.group(1)).strip() for m in _HTML_SHEET_HEADING_RE.finditer(html)
    ]
    by_name = {name: widths for name, widths in sheets}
    plan: list[list[int]] = []
    for i, ncols in enumerate(counts):
        widths: list[int] | None = None
        if i < len(headings):
            widths = by_name.get(headings[i])
        if widths is None and not headings and i < len(sheets):
            widths = sheets[i][1]  # heading format changed → positional fallback
        if widths and 0 < ncols <= len(widths):
            plan.append(widths[:ncols])
        else:
            plan.append([])
    return plan


class _OfficeHtmlSanitizer(HTMLParser):
    """Whitelist sanitizer for LibreOffice-generated spreadsheet HTML.

    Cell text is HTML-escaped by LibreOffice, but formula cells can emit live
    anchors (``=HYPERLINK("http://…")``) and a hostile workbook could carry
    other active markup, so the preview document is rebuilt from scratch:
    only whitelisted elements/attributes survive, external URLs and ``on*``
    handlers are dropped, local chart images are inlined as data URIs, and
    the LibreOffice stylesheet is kept only when it contains no off-disk
    ``url()``/``@import`` references.
    """

    def __init__(self, assets_dir: Path, table_widths: list[list[int]] | None = None):
        super().__init__(convert_charrefs=True)
        self._assets_dir = assets_dir
        self._table_widths = table_widths or []
        self._table_index = 0
        # LibreOffice emits one <colgroup> per width run; after injecting the
        # workbook widths every original colgroup/col is dropped until the
        # first data row (or the end of the table).
        self._skip_colgroups = False
        self._open_colgroup = 0
        self._body: list[str] = []
        self._styles: list[str] = []
        self._in_head = False
        self._in_style = False
        self._style_buf: list[str] = []
        self._suppress: list[str] = []
        self._data_uri_cache: dict[str, str | None] = {}
        self._inlined_asset_bytes = 0

    # -- helpers ----------------------------------------------------------
    def _local_asset_data_uri(self, src: str) -> str | None:
        # LibreOffice percent-encodes non-ASCII chart image names on Linux
        # (LO 7.3 prod), so decode before resolving the flat export file.
        raw = unquote(src.strip())
        name = os.path.basename(raw)
        if not name:
            return None
        if name in self._data_uri_cache:
            return self._data_uri_cache[name]
        result: str | None = None
        path = self._assets_dir / name
        if path.suffix.lower() in _HTML_ASSET_EXT and path.is_file():
            try:
                size = path.stat().st_size
                if size <= _HTML_ASSET_MAX_BYTES:
                    if self._inlined_asset_bytes + size > _HTML_ASSET_TOTAL_MAX_BYTES:
                        raise OfficePreviewFailed("HTML preview assets exceed size budget")
                    mime = _HTML_ASSET_MIME.get(path.suffix.lower(), "application/octet-stream")
                    result = f"data:{mime};base64,{base64.b64encode(path.read_bytes()).decode('ascii')}"
                    self._inlined_asset_bytes += size
            except OSError:
                result = None
        self._data_uri_cache[name] = result
        return result

    def _render_attrs(self, tag: str, attrs: list[tuple[str, str | None]]) -> str:
        parts: list[str] = []
        for raw_name, raw_value in attrs:
            name = (raw_name or "").lower()
            value = "" if raw_value is None else raw_value
            if not name or name.startswith("on") or name.startswith("data-"):
                continue
            if name == "style":
                if any(token in value.lower() for token in _HTML_CSS_FORBIDDEN):
                    continue
            elif name == "href":
                if not value.strip().startswith("#"):
                    continue
            elif name == "src":
                if not value.strip().lower().startswith("data:image/"):
                    continue
            elif name not in _HTML_ALLOWED_ATTRS:
                continue
            parts.append(f'{name}="{html_lib.escape(value, quote=True)}"')
        return (" " + " ".join(parts)) if parts else ""

    # -- HTMLParser hooks -------------------------------------------------
    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag in _HTML_SUPPRESS_TAGS:
            self._suppress.append(tag)
            return
        if self._suppress:
            return
        if tag == "head":
            self._in_head = True
            return
        if tag == "style":
            self._in_style = True
            self._style_buf = []
            return
        if tag == "table":
            widths = (
                self._table_widths[self._table_index]
                if self._table_index < len(self._table_widths)
                else []
            )
            self._table_index += 1
            if widths:
                total = sum(widths)
                cols = "".join(f'<col style="width:{w}px">' for w in widths)
                self._body.append(f'<table style="table-layout:fixed;width:{total}px">')
                self._body.append(f"<colgroup>{cols}</colgroup>")
                self._skip_colgroups = True  # drop stale LibreOffice colgroups
            else:
                self._body.append(f"<table{self._render_attrs(tag, attrs)}>")
            return
        if tag == "tr":
            self._skip_colgroups = False  # colgroups always precede the rows
        if tag in ("colgroup", "col") and self._skip_colgroups:
            return
        if tag in _HTML_SKELETON_TAGS or tag not in _HTML_ALLOWED_TAGS:
            return
        if tag == "img":
            rendered = self._render_attrs(tag, attrs)
            src_attr = None
            for raw_name, raw_value in attrs:
                if (raw_name or "").lower() == "src":
                    src_attr = "" if raw_value is None else raw_value.strip()
                    break
            if src_attr is None:
                return
            if not src_attr.lower().startswith("data:image/"):
                inlined = self._local_asset_data_uri(src_attr)
                if not inlined:
                    return  # unresolvable/external image → drop the element
                rendered = re.sub(
                    r'src="[^"]*"', f'src="{inlined}"', rendered, count=1
                ) if 'src="' in rendered else f' src="{inlined}"'
            self._body.append(f"<img{rendered}>")
            return
        if tag == "colgroup":
            self._open_colgroup += 1  # auto-layout table keeps its colgroup
        self._body.append(f"<{tag}{self._render_attrs(tag, attrs)}>")

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag in _HTML_SUPPRESS_TAGS:
            return
        if tag in ("colgroup", "col") and self._skip_colgroups:
            return
        if tag == "img":
            self.handle_starttag(tag, attrs)
            return
        if tag in _HTML_ALLOWED_TAGS and tag not in _HTML_SKELETON_TAGS:
            self._body.append(f"<{tag}{self._render_attrs(tag, attrs)}>")

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if self._suppress:
            if tag == self._suppress[-1]:
                self._suppress.pop()
            return
        if tag == "head":
            self._in_head = False
            return
        if tag == "style":
            self._in_style = False
            css = "".join(self._style_buf)
            if css and not any(token in css.lower() for token in _HTML_CSS_FORBIDDEN):
                self._styles.append(css)
            self._style_buf = []
            return
        if tag == "colgroup":
            if self._open_colgroup > 0:
                self._open_colgroup -= 1
                self._body.append("</colgroup>")
            return
        if tag == "table":
            self._skip_colgroups = False
        if (
            tag in _HTML_ALLOWED_TAGS
            and tag not in _HTML_VOID_TAGS
            and tag not in _HTML_SKELETON_TAGS
        ):
            self._body.append(f"</{tag}>")

    def handle_data(self, data: str) -> None:
        if self._suppress:
            return
        if self._in_style:
            self._style_buf.append(data)
            return
        if self._in_head:
            return
        self._body.append(html_lib.escape(data, quote=False))

    def handle_comment(self, _data: str) -> None:
        return

    def handle_decl(self, _decl: str) -> None:
        return

    def render(self) -> str:
        styles = "".join(self._styles)
        doc = (
            f"<!DOCTYPE html><html {_HTML_PREVIEW_MARKER}><head><meta charset=\"utf-8\">"
            f"<style>{styles}</style>"
            f"<style>{_HTML_PREVIEW_CSS}</style>"
            "</head><body>" + "".join(self._body) + "</body></html>"
        )
        if len(doc.encode("utf-8")) > _HTML_MAX_BYTES:
            raise OfficePreviewFailed(
                f"HTML preview exceeds {_HTML_MAX_BYTES // (1024 * 1024)}MB limit"
            )
        return doc


def sanitize_office_html(
    html: str,
    assets_dir: str | Path,
    table_widths: list[list[int]] | None = None,
) -> str:
    """Sanitize one LibreOffice HTML export and inline its local images.

    ``table_widths`` (optional) re-injects the source workbook's column widths
    per HTML table so wide tables keep their natural width and scroll
    horizontally; unresolvable tables keep the browser's auto layout.
    """
    if len(html) > _HTML_MAX_RAW_CHARS:
        raise OfficePreviewFailed("HTML preview source exceeds size limit")
    if table_widths:
        counts = _html_table_column_counts(html)
        aligned: list[list[int]] = []
        for i, widths in enumerate(table_widths):
            ncols = counts[i] if i < len(counts) else 0
            aligned.append(widths[:ncols] if 0 < ncols <= len(widths) else [])
        table_widths = aligned
    parser = _OfficeHtmlSanitizer(Path(assets_dir), table_widths)
    parser.feed(html)
    parser.close()
    return parser.render()


def _find_soffice(preferred: str) -> str | None:
    if not preferred:
        preferred = "soffice"
    if os.sep in preferred or (os.altsep and os.altsep in preferred):
        if os.path.isfile(preferred) and os.access(preferred, os.X_OK):
            return preferred
        return None
    found = shutil.which(preferred)
    if found:
        return found
    if preferred == "soffice":
        for candidate in _COMMON_SOFFICE_PATHS:
            if os.path.isfile(candidate) and os.access(candidate, os.X_OK):
                return candidate
    return None


async def ensure_soffice(preferred: str = "soffice") -> str | None:
    """Return an executable soffice path, or ``None`` when unavailable."""
    import asyncio

    return await asyncio.to_thread(_find_soffice, preferred)


def _is_valid_pdf(path: Path) -> bool:
    """Cache entries are only trusted after a magic-byte check (A4.9 minor:
    a poisoned shared-temp cache entry must not be served as a preview)."""
    try:
        with open(path, "rb") as handle:
            return handle.read(5) == b"%PDF-"
    except OSError:
        return False


def _cache_key(source_id: str, mtime_ns: int, size: int, ext: str) -> str:
    raw = f"v{_CACHE_VERSION}:{source_id}:{mtime_ns}:{size}:{ext.lower()}"
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:32]


def convert_target(ext: str) -> str:
    """LibreOffice ``--convert-to`` argument for a source extension.

    Calc single-page export requires LibreOffice >= 7.4 for JSON filter
    options; older builds ignore them (conversion still succeeds, paginated).
    Hidden sheets are included in single-page mode by LibreOffice design —
    acceptable here (the workspace preview shows the whole workbook).
    """
    if ext.lower() in _CALC_SINGLE_PAGE_EXT:
        return 'pdf:calc_pdf_Export:{"SinglePageSheets":{"type":"boolean","value":"true"}}'
    return "pdf"


_SHEET_REF_RE = re.compile(rb'<dimension[^>]*ref="([^"]+)"[^>]*/?>')
_SHEET_ROW_RE = re.compile(rb'<row[^>]* r="(\d+)"')
_CELL_ROW_RE = re.compile(rb'([A-Za-z]{1,3})\$?(\d+)$')


def _sheet_max_row(xml: bytes) -> int:
    """Best-effort max row of a worksheet XML (0 when undeterminable)."""
    match = _SHEET_REF_RE.search(xml)
    if match:
        ref = match.group(1).decode("ascii", errors="ignore").replace("$", "").strip("'\"")
        last = ref.split(":")[-1]
        cell = _CELL_ROW_RE.search(last.encode("ascii", errors="ignore"))
        if cell:
            return int(cell.group(2))
    rows = _SHEET_ROW_RE.findall(xml)
    return max((int(r) for r in rows), default=0)


def _xlsx_needs_pagination(src: Path) -> bool:
    """True when any sheet exceeds the fit-to-one-page row guard."""
    try:
        with zipfile.ZipFile(src) as archive:
            for name in archive.namelist():
                if re.match(r"xl/worksheets/sheet\d+\.xml$", name):
                    if _sheet_max_row(archive.read(name)) > _XLSX_FIT_MAX_ROWS:
                        return True
    except (OSError, zipfile.BadZipFile, ValueError, RuntimeError, NotImplementedError):
        return False
    return False


def _patch_xlsx_print_setup(src: Path, workdir: Path | None = None) -> Path:
    """Copy an .xlsx spool with every sheet set to fit-to-one-page printing.

    Pure zip rewrite — charts/drawings/styles are copied byte-for-byte (an
    openpyxl round-trip would drop charts). Returns the patched path, or
    *src* unchanged when the file is not a patchable zip or already fits.
    """
    patched = (workdir or src.parent) / (src.stem + ".wt_fit.xlsx")
    try:
        changed = False
        with zipfile.ZipFile(src) as zin, zipfile.ZipFile(patched, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if re.match(r"xl/worksheets/sheet\d+\.xml$", item.filename):
                    xml = data.decode("utf-8")
                    new_xml = _patch_sheet_fit_to_page(xml)
                    if new_xml != xml:
                        # never accept a rewrite that isn't well-formed XML —
                        # a corrupt sheet is silently dropped by LibreOffice
                        # while the wrong PDF gets cached (A4.9 fix8 I-1)
                        try:
                            ET.fromstring(new_xml)
                        except ET.ParseError:
                            patched.unlink(missing_ok=True)
                            return src
                        changed = True
                    data = new_xml.encode("utf-8")
                zout.writestr(item, data)
        if not changed:
            patched.unlink(missing_ok=True)
            return src
        return patched
    except (OSError, zipfile.BadZipFile, ValueError, RuntimeError, NotImplementedError):
        try:
            patched.unlink(missing_ok=True)
        except OSError:
            pass
        return src


def _patch_sheet_fit_to_page(xml: str) -> str:
    fit = '<pageSetUpPr fitToPage="1"/>'
    if "<pageSetUpPr" in xml:
        # replace the WHOLE element (self-closing or paired) — a partial
        # replace would orphan the closing tag and corrupt the XML (A4.9
        # fix8 Important-1)
        xml = re.sub(
            r"<pageSetUpPr\b[^>]*?(?:/>|>.*?</pageSetUpPr>)",
            fit,
            xml,
            count=1,
            flags=re.DOTALL,
        )
    else:
        self_close = re.search(r"<sheetPr\s*/>", xml)
        open_tag = re.search(r"<sheetPr(?:\s[^>]*)?>", xml)
        if self_close:
            xml = xml.replace(self_close.group(0), f"<sheetPr>{fit}</sheetPr>", 1)
        elif open_tag:
            # pageSetUpPr is last in CT_SheetPr — insert before </sheetPr>
            if "</sheetPr>" in xml:
                xml = xml.replace("</sheetPr>", fit + "</sheetPr>", 1)
            else:
                xml = xml.replace(open_tag.group(0), open_tag.group(0) + fit, 1)
        else:
            # `<sheetProtection ...>` must not be mistaken for sheetPr
            # (A4.9 fix8 Important-2): anchor on the real opening tag only
            xml = re.sub(
                r"(<worksheet\b[^>]*>)",
                r"\1<sheetPr>" + fit + "</sheetPr>",
                xml,
                count=1,
            )
    if "<pageSetup" in xml:
        match = re.search(r"<pageSetup\b[^>]*?/?>", xml)
        tag = match.group(0)
        new_tag = tag
        if 'fitToWidth="' in new_tag:
            new_tag = re.sub(r'fitToWidth="[^"]*"', 'fitToWidth="1"', new_tag)
        else:
            new_tag = new_tag[:-2] + ' fitToWidth="1"/>' if new_tag.endswith("/>") else new_tag[:-1] + ' fitToWidth="1">'
        if 'fitToHeight="' in new_tag:
            new_tag = re.sub(r'fitToHeight="[^"]*"', 'fitToHeight="1"', new_tag)
        else:
            new_tag = new_tag[:-2] + ' fitToHeight="1"/>' if new_tag.endswith("/>") else new_tag[:-1] + ' fitToHeight="1">'
        xml = xml.replace(tag, new_tag, 1)
    else:
        xml = xml.replace(
            "</worksheet>",
            '<pageSetup fitToWidth="1" fitToHeight="1"/></worksheet>',
        )
    return xml


def _pdf_page_sizes(raw: bytes) -> list[tuple[float, float]]:
    """(width, height) per page from raw PDF bytes; [] when unparseable."""
    sizes: list[tuple[float, float]] = []
    for match in _MEDIABOX_RE.finditer(raw):
        try:
            nums = [float(x) for x in match.group(1).split()]
        except ValueError:
            continue
        if len(nums) == 4:
            sizes.append((abs(nums[2] - nums[0]), abs(nums[3] - nums[1])))
    return sizes


def _needs_paginated_fallback(pdf_path: Path) -> bool:
    try:
        raw = pdf_path.read_bytes()
    except OSError:
        return False
    sizes = _pdf_page_sizes(raw)
    return bool(sizes) and max(height for _w, height in sizes) > _SINGLE_PAGE_MAX_PT


def _prune_stale_spools(cache_dir: Path, max_age_seconds: int = 3600) -> None:
    """Remove abandoned spool copies (crash between spool and unlink)."""
    import time

    cutoff = time.time() - max_age_seconds
    try:
        candidates = list(cache_dir.glob("wt_office_src_*"))
    except OSError:
        return
    for path in candidates:
        try:
            if path.is_file() and path.stat().st_mtime < cutoff:
                path.unlink()
        except OSError:
            continue


def _prune_cache(cache_dir: Path, max_cache_mb: int) -> None:
    max_bytes = max(0, int(max_cache_mb)) * 1024 * 1024
    if max_bytes <= 0:
        return
    try:
        files = [
            p
            for p in cache_dir.iterdir()
            if p.is_file() and p.suffix.lower() in (".pdf", ".html")
        ]
    except OSError:
        return
    total = sum(p.stat().st_size for p in files if p.exists())
    if total <= max_bytes:
        return
    for path in sorted(files, key=lambda p: p.stat().st_mtime):
        if total <= max_bytes:
            break
        try:
            size = path.stat().st_size
            path.unlink()
            total -= size
        except OSError:
            continue


def _convert_once(
    soffice: str,
    src: Path,
    profile: Path,
    target: str,
    timeout: int,
    out_ext: str = "pdf",
) -> Path:
    """One soffice run into a private outdir; returns the produced file."""
    outdir = profile.parent / "out"
    outdir.mkdir(exist_ok=True)
    for leftover in outdir.iterdir():
        if leftover.is_file():
            leftover.unlink()
    cmd = [
        soffice,
        f"-env:UserInstallation=file://{profile}",
        "--headless",
        "--norestore",
        "--convert-to",
        target,
        "--outdir",
        str(outdir),
        str(src),
    ]
    proc = subprocess.Popen(
        cmd,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        start_new_session=True,
    )
    try:
        _out, err = proc.communicate(timeout=max(5, int(timeout)))
    except subprocess.TimeoutExpired:
        try:
            os.killpg(os.getpgid(proc.pid), signal.SIGKILL)
        except OSError:
            proc.kill()
        try:
            proc.wait(timeout=10)
        except subprocess.TimeoutExpired:
            pass
        raise OfficePreviewFailed(f"conversion timed out after {timeout}s")
    if proc.returncode != 0:
        detail = (err or b"").decode("utf-8", errors="replace")[:300]
        raise OfficePreviewFailed(f"soffice exited {proc.returncode}: {detail}")
    produced = [
        p for p in sorted(outdir.iterdir()) if p.is_file() and p.suffix.lower() == f".{out_ext}"
    ]
    if not produced:
        detail = (err or b"").decode("utf-8", errors="replace")[:300]
        raise OfficePreviewFailed(f"no {out_ext.upper()} produced: {detail}")
    return produced[0]


def _run_soffice(
    soffice: str,
    src: Path,
    dest: Path,
    timeout: int,
    target: str | None = None,
    transform: "Callable[[Path], Path] | None" = None,
) -> None:
    """Blocking conversion; runs in a worker thread.

    ``target`` overrides the LibreOffice ``--convert-to`` argument (used by
    the HTML preview path); when it is ``None`` the PDF pipeline below runs
    unchanged. ``transform`` post-processes the produced artifact inside the
    conversion tempdir (chart images etc. are still on disk there).
    """
    if not (os.path.isfile(soffice) and os.access(soffice, os.X_OK)):
        resolved = shutil.which(soffice)
        if not resolved:
            raise OfficePreviewUnavailable(f"soffice not found: {soffice}")
        soffice = resolved

    with tempfile.TemporaryDirectory(prefix="wt_office_preview_") as tmp:
        profile = Path(tmp) / "profile"
        convert_src = src
        chosen = target if target is not None else convert_target(src.suffix)
        if (
            target is None
            and src.suffix.lower() == _XLSX_EXT
            and not _xlsx_needs_pagination(src)
        ):
            # Version-independent single-sheet pages: patch print setup, then
            # plain PDF export (works on LibreOffice 7.3 too — prod box).
            convert_src = _patch_xlsx_print_setup(src, Path(tmp))
            if convert_src != src:
                chosen = "pdf"
        out_ext = chosen.split(":")[0]
        produced = _convert_once(soffice, convert_src, profile, chosen, timeout, out_ext)

        # Single-page export can flatten a huge sheet into one unrenderably
        # tall page (browser canvas limits → silent blank preview). Fall back
        # to the paginated export for those (A4.9 I-2).
        if out_ext == "pdf" and chosen != "pdf" and _needs_paginated_fallback(produced):
            produced = _convert_once(soffice, src, Path(tmp) / "profile2", "pdf", timeout)
        if transform is not None:
            produced = transform(produced)

        dest.parent.mkdir(parents=True, exist_ok=True)
        staging = dest.with_name(dest.name + ".tmp")
        shutil.move(str(produced), str(staging))
        os.replace(staging, dest)


async def convert_to_pdf(
    src: str | Path,
    *,
    soffice: str,
    cache_dir: str | Path,
    timeout: int = 60,
    max_cache_mb: int = 512,
    cache_identity: str | None = None,
    cache_stat: tuple[int, int] | None = None,
) -> Path:
    """Convert *src* to a cached PDF and return the PDF path.

    ``cache_identity``/``cache_stat`` let callers key the cache on the
    *original workspace source* (path + fstat mtime/size) when *src* is a
    random-named spool copy — otherwise every request would miss the cache.
    Raises :class:`OfficePreviewUnavailable` when soffice is missing and
    :class:`OfficePreviewFailed` when conversion fails or times out.
    """
    import asyncio

    src_path = Path(src)
    if not is_supported(src_path.name):
        raise OfficePreviewFailed(f"unsupported office extension: {src_path.suffix}")
    try:
        stat_result = src_path.stat()
    except OSError as exc:
        raise OfficePreviewFailed(f"cannot stat source: {exc}") from exc
    if not src_path.is_file():
        raise OfficePreviewFailed("source is not a regular file")

    cache_path = Path(cache_dir)
    try:
        cache_path.mkdir(parents=True, exist_ok=True)
        os.chmod(cache_path, 0o700)
    except OSError:
        pass
    identity = cache_identity or str(src_path.resolve())
    mtime_ns, size = cache_stat if cache_stat else (stat_result.st_mtime_ns, stat_result.st_size)
    dest = cache_path / f"{_cache_key(identity, mtime_ns, size, src_path.suffix)}.pdf"
    _prune_stale_spools(cache_path)
    if _is_valid_pdf(dest):
        return dest
    if dest.exists():
        try:
            dest.unlink()
        except OSError:
            pass

    async with _conversion_semaphore():
        if _is_valid_pdf(dest):
            return dest
        await asyncio.to_thread(_run_soffice, soffice, src_path, dest, timeout)
        await asyncio.to_thread(_prune_cache, cache_path, max_cache_mb)
        return dest


async def convert_to_html(
    src: str | Path,
    *,
    soffice: str,
    cache_dir: str | Path,
    timeout: int = 60,
    max_cache_mb: int = 512,
    cache_identity: str | None = None,
    cache_stat: tuple[int, int] | None = None,
) -> Path:
    """Convert a spreadsheet to a cached, sanitized, self-contained HTML file.

    LibreOffice's Calc HTML export renders every sheet as one natural-width
    table (charts become images); :func:`sanitize_office_html` inlines those
    images as data URIs, strips active/external markup and injects the
    wrapping/scroll stylesheet. The artifact is served to a sandboxed iframe,
    so no LibreOffice- or workbook-controlled script ever runs.

    Cache semantics and error classes mirror :func:`convert_to_pdf`.
    """
    import asyncio

    src_path = Path(src)
    if not is_spreadsheet(src_path.name):
        raise OfficePreviewFailed(f"not a spreadsheet: {src_path.suffix}")
    try:
        stat_result = src_path.stat()
    except OSError as exc:
        raise OfficePreviewFailed(f"cannot stat source: {exc}") from exc
    if not src_path.is_file():
        raise OfficePreviewFailed("source is not a regular file")

    cache_path = Path(cache_dir)
    try:
        cache_path.mkdir(parents=True, exist_ok=True)
        os.chmod(cache_path, 0o700)
    except OSError:
        pass
    identity = cache_identity or str(src_path.resolve())
    mtime_ns, size = cache_stat if cache_stat else (stat_result.st_mtime_ns, stat_result.st_size)
    dest = cache_path / f"{_cache_key(identity, mtime_ns, size, src_path.suffix)}.html"
    _prune_stale_spools(cache_path)
    if _is_valid_html(dest):
        return dest
    if dest.exists():
        try:
            dest.unlink()
        except OSError:
            pass

    def _transform(produced: Path) -> Path:
        if produced.stat().st_size > _HTML_MAX_RAW_CHARS:
            raise OfficePreviewFailed("HTML preview source exceeds size limit")
        raw = produced.read_text(encoding="utf-8", errors="replace")
        widths: list[list[int]] | None = None
        sheets = _sheet_column_widths(src_path)
        if sheets:
            widths = _html_table_width_plan(raw, sheets)
        sanitized = sanitize_office_html(raw, produced.parent, widths)
        clean = produced.with_name(produced.stem + ".wt_preview.html")
        clean.write_text(sanitized, encoding="utf-8")
        return clean

    async with _conversion_semaphore():
        if _is_valid_html(dest):
            return dest
        await asyncio.to_thread(
            _run_soffice, soffice, src_path, dest, timeout, "html", _transform
        )
        await asyncio.to_thread(_prune_cache, cache_path, max_cache_mb)
        return dest
